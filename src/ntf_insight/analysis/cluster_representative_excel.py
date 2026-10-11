"""Excel-Arbeitsliste für reale FIN-Repräsentanten der lokalen Clusteranalyse."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ntf_insight.analysis.autoexport_fin_patterns import FinPatternResult


def _safe(value: object) -> str:
    """Prevent Excel formula injection from imported, untrusted text."""
    text = str(value if value is not None else "")
    if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + text
    return text


def export_cluster_representatives(
    path: str | Path, result: FinPatternResult,
    *, filter_description: str = "",
) -> int:
    """One representative FIN per cluster, plus a manual download checklist."""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "FIN-Repräsentanten"
    headers = (
        "Cluster", "Repräsentative FIN", "FINs im Cluster",
        "FINs mit QEC", "Repräsentant mit QEC", "QEC-Fallkennungen",
        "DTC-Merkmale", "Gemeinsame ECU:DTCs", "Median Ähnlichkeit",
        "Datenherkunft", "Befundungsbericht herunterladen", "Bericht geprüft",
        "Notizen",
    )
    ws.append(headers)
    for group in sorted(result.groups, key=lambda group: group.group_id):
        fin = group.representative_fin
        common = "; ".join(
            f"{ecu}:{dtc} ({count}/{group.fin_count})"
            for ecu, dtc, count, _ in group.common_faults
        )
        ws.append((
            f"C{group.group_id}", _safe(fin), group.fin_count,
            group.qec_fin_count, "Ja" if group.representative_has_qec else "Nein",
            _safe("; ".join(result.qec_links.get(fin, ()))),
            result.code_counts.get(fin, 0), _safe(common),
            group.median_similarity,
            _safe((result.sources_by_fin or {}).get(fin, "")),
            "Offen", "Offen", "",
        ))
        ws.cell(ws.max_row, 9).number_format = "0.0%"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:M{max(1, ws.max_row)}"
    ws.sheet_view.showGridLines = False
    widths = (13, 24, 18, 17, 24, 36, 17, 65, 21, 23, 33, 18, 34)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="22252A")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 31
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="333333")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in (11, 12, 13):
            row[col - 1].font = Font(name="Arial", size=10, color="2456A6")
    info = workbook.create_sheet("Analyseinformationen")
    info.append(("Feld", "Wert"))
    for key, value in (
        ("Verfahren", result.method),
        ("Analysierte FINs", result.total_fins),
        ("FINs mit DTC-Merkmalen", result.dtc_fins),
        ("Cluster", len(result.groups)),
        ("DTC-Filter", filter_description or "Keine"),
        ("Hinweis", "Diagnoseähnlichkeit beweist weder einen Pumpendefekt noch Lieferantenverschulden."),
        ("Datenschutz", "Datei enthält echte FINs; nur in freigegebenen lokalen Ablagen speichern."),
    ):
        info.append((_safe(key), _safe(value)))
    info.column_dimensions["A"].width = 27
    info.column_dimensions["B"].width = 110
    for c in info[1]:
        c.fill = PatternFill("solid", fgColor="22252A")
        c.font = Font(bold=True, color="FFFFFF")
    workbook.save(path)
    return len(result.groups)
