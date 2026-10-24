"""Erkennung und Ein-Datei-Regel der QEC-generierten FIN-Liste.

Das Excel-Blatt "Schadteile QEC" mit FIN ist die *primäre* Kohorte.
Eine ergänzende SSL-Endbefundliste ("Schadteile" ohne FIN) ist erlaubt;
sie erzeugt keine zweite FIN-Kohorte.
"""
from __future__ import annotations

from pathlib import Path

from ntf_insight.ingestion.artifacts import SourceArtifact, sha256_file


def primary_qec_sheet(workbook) -> str | None:
    """Identify real QEC FIN sheets by their header, never by an export name.

    Supports original Schadteile QEC and generated filtered QEC workbooks.
    """
    required = {"Referenznr.", "FIN", "Rep.-Datum"}
    matching = []
    for sheet in workbook:
        headers = next(
            sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()
        )
        if required <= {
            str(value).strip() for value in headers if value is not None
        }:
            matching.append(sheet.title)
    if "Schadteile QEC" in matching:
        return "Schadteile QEC"
    if len(matching) > 1:
        raise ValueError("Mehrere QEC-FIN-Blätter erkannt. Excel-Datei prüfen.")
    return matching[0] if matching else None


def is_primary_qec_workbook(path: Path) -> bool:
    """True only for a primary QEC export, never for an SSL-only workbook."""
    path = Path(path)
    if path.suffix.casefold() != ".xlsx" or not path.is_file():
        return False
    from openpyxl import load_workbook

    workbook = load_workbook(
        path, read_only=True, data_only=True, keep_links=False,
    )
    try:
        return primary_qec_sheet(workbook) is not None
    finally:
        workbook.close()


def active_qec_lists(artifacts: list[SourceArtifact]) -> list[SourceArtifact]:
    """Identify all primary workbooks, including databases imported before this rule."""
    return [
        artifact for artifact in artifacts
        if is_primary_qec_workbook(Path(artifact.archived_path))
    ]


def _part_keys(path: Path) -> set[str]:
    """Read the actual QEC case part numbers, not a guessed filename."""
    from ntf_insight.ingestion.real_source_parsers import QECWorkbookParser

    parsed = QECWorkbookParser(sha256_file(path)).parse(path).cases
    values = {
        "".join(c for c in str(case.qec_part_number or "").upper() if c.isalnum())
        for case in parsed
    }
    if not values or "" in values:
        raise ValueError(
            "QEC-FIN-Liste enthält Fälle ohne gültige Sachnummer. "
            "Eine projektbezogene Zuordnung ist damit nicht sicher."
        )
    return values


def validate_single_qec_list(
    incoming_paths: list[Path],
    archived: list[SourceArtifact],
) -> None:
    """At most one *different* QEC primary source per part number.

    Multiple projects may coexist, but overlapping QEC cohorts require an
    explicit source replacement rather than silently overriding decisions.
    """
    sources = {
        artifact.sha256: Path(artifact.archived_path)
        for artifact in active_qec_lists(archived)
    }
    for path in incoming_paths:
        if is_primary_qec_workbook(path):
            sources[sha256_file(path)] = Path(path)
    ownership: dict[str, str] = {}
    for sha, path in sorted(sources.items()):
        for part in _part_keys(path):
            old_sha = ownership.setdefault(part, sha)
            if old_sha != sha:
                raise ValueError(
                    "Pro Sachnummer darf nur eine QEC-generierte FIN-Liste "
                    "aktiv sein. Zwei verschiedene Quelldateien enthalten "
                    f"dieselbe Sachnummer {part}. Bitte die alte Liste "
                    "vor einem Neuimport gezielt ersetzen."
                )
