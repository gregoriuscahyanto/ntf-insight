"""Conservative content-based classification of supported local source files.

Do not rely on proprietary filenames, and never execute or upload their contents.
"""
from __future__ import annotations

import csv
from pathlib import Path

from ntf_insight.ingestion.artifacts import SourceType
from ntf_insight.ingestion.qec_primary import primary_qec_sheet


def detect_source_type(path: Path) -> SourceType:
    path = Path(path)
    ext = path.suffix.casefold()

    if ext == ".json":
        from ntf_insight.ingestion.qec_json import QECJsonParser
        from ntf_insight.ingestion.tips_json import TipsJsonParser
        for parser, source_type in (
            (QECJsonParser, SourceType.QEC_REPORT),
            (TipsJsonParser, SourceType.TIPS_CASE),
        ):
            try:
                parser().parse(path)
            except (OSError, UnicodeError, ValueError, TypeError):
                continue
            return source_type
        return SourceType.OTHER

    if ext == ".smr-d":
        return SourceType.SMR_D
    if ext == ".cdd":
        with path.open("rb") as source:
            prefix = source.read(2048)
        return (
            SourceType.DIAGNOSTIC_DESCRIPTION
            if b"<CANDELA" in prefix
            else SourceType.OTHER
        )
    if ext == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
            for line in source:
                if line.startswith(("Original query:", "Query Id:")) or not line.strip():
                    continue
                names = next(csv.reader([line], delimiter=";"))
                headers = set(names)
                if {"Fin", "CaseId", "CreationDate", "InitialWorkshopMessageCustomerComplaint"} <= headers:
                    return SourceType.TIPS_CASE
                if {"DiagStartDt", "EcuFaultNm", "Fin"} <= headers:
                    if {
                        "EcuNm", "CntXDDiagnosisHeaderSid", "LastRepairDt", "C1"
                    } <= headers:
                        return SourceType.TP_DASHBOARD
                    return SourceType.SHORT_TEST
                if {"RepairDt", "Fin", "PartLetterCd"} <= headers:
                    return SourceType.WORKSHOP_HISTORY
                if names == ["C1"]:
                    return SourceType.AUTOEXPORT_SUMMARY
                break
        return SourceType.OTHER
    if ext == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if primary_qec_sheet(workbook) is not None:
                return SourceType.VEHICLE_CASE_LIST
            # QEC export variant without VIN: SSL result can be joined only
            # to an ALREADY known QEC case with the exact reference number.
            # Recognizing the file never creates a case or guesses a FIN.
            if "Schadteile" in workbook.sheetnames:
                header = next(
                    workbook["Schadteile"].iter_rows(
                        min_row=1, max_row=1, values_only=True
                    ),
                    (),
                )
                if {
                    "Referenznr.", "SSL Endbefund",
                    "SSL-Endbefund Benennung",
                } <= set(header):
                    return SourceType.VEHICLE_CASE_LIST
        finally:
            workbook.close()
        return SourceType.OTHER
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path, strict=False)
        if not reader.pages:
            return SourceType.OTHER
        text = (reader.pages[0].extract_text() or "")[:15000]
        if "DIAG START" in text and "QuickTest" in text and "ECUS (" in text:
            return SourceType.SHORT_TEST
        if "Part History List" in text and ("Datum zavedení" in text or "Q01" in text):
            return SourceType.PART_HISTORY
        if "ANALYSIS" in text and "Complaint" in text:
            return SourceType.SUPPLIER_FINDING
        return SourceType.OTHER
    return SourceType.OTHER
