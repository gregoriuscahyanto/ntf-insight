"""Structured, local-only reference parsers for Candela, part history and QEC.

CDD and Part History describe hardware/product knowledge, NOT a FIN-specific
DTC occurrence. Neither can fabricate repair events or SMR-D trigger rules.
"""
from __future__ import annotations

import csv
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterator

from pydantic import BaseModel, Field


class KnowledgeResult(BaseModel):
    kind: str
    metadata: dict = Field(default_factory=dict)
    records: list[dict] = Field(default_factory=list)
    issues: list[str] = Field(default_factory=list)


def _safe_cell(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (int, float, bool, str)):
        return value
    return str(value)


class QECDetailsParser:
    """Preserve every sheet/column independently of the canonical QEC cases."""

    name = "qec_details"
    version = "1.0.0"

    def parse(self, path: Path) -> KnowledgeResult:
        from openpyxl import load_workbook

        if path.stat().st_size > 50_000_000:
            raise ValueError("QEC Excel ist zu groß.")
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        result = KnowledgeResult(kind="qec_details")
        try:
            if len(workbook.sheetnames) > 30:
                raise ValueError("QEC enthält unerwartet viele Tabellenblätter.")
            count = 0
            summaries = {}
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, None)
                if not headers:
                    continue
                # Preserve duplicate column labels without discarding values.
                names = []
                seen = {}
                for index, name in enumerate(headers, start=1):
                    key = str(name).strip() if name is not None else f"Spalte_{index}"
                    seen[key] = seen.get(key, 0) + 1
                    names.append(key if seen[key] == 1 else f"{key}__{seen[key]}")
                local_count = 0
                for row_number, values in enumerate(rows, start=2):
                    if not any(value is not None for value in values):
                        continue
                    count += 1
                    if count > 50_000:
                        raise ValueError("QEC-Limit von 50.000 Datenzeilen erreicht.")
                    data = {
                        name: _safe_cell(value)
                        for name, value in zip(names, values)
                    }
                    result.records.append({
                        "category": "qec_row",
                        "record_number": count,
                        "source_location": f"{sheet.title}!{row_number}",
                        "fin": str(data.get("FIN") or "").strip() or None,
                        "reference": str(data.get("Referenznr.") or "").strip() or None,
                        "label": str(data.get("SSL-Endbefund Benennung") or data.get("Status") or ""),
                        "data": data,
                    })
                    local_count += 1
                summaries[sheet.title] = {"rows": local_count, "columns": names}
            result.metadata = {"sheets": summaries, "total_rows": count}
            if count == 0:
                result.issues.append("Keine QEC-Zeilen extrahiert.")
        finally:
            workbook.close()
        return result


class CandelaParser:
    """Extract Candela diagnostic services using streaming, XXE-disabled XML.

    This is *not* an SMR-D trigger parser. The supplied Candela file contains
    service/instance definitions but no validated mapping of Set/Enable/Healing
    rules for individual field DTCs.
    """

    name = "candela_services"
    version = "1.0.0"

    def parse(self, path: Path) -> KnowledgeResult:
        from lxml import etree

        if path.stat().st_size > 160_000_000:
            raise ValueError("CDD überschreitet die Größenbegrenzung (160 MB).")
        with path.open("rb") as stream:
            header = stream.read(65536)
        if b"<CANDELA" not in header or b"<!ENTITY" in header.upper():
            raise ValueError("Ungültiges oder unsicheres CANDELA-Dokument.")

        out = KnowledgeResult(
            kind="candela_services",
            metadata={"standard": "CANDELA", "triggers_validated": False},
        )
        class_qualifiers: dict[str, str] = {}
        inside = 0
        service_index = 0
        with path.open("rb") as stream:
            for action, node in etree.iterparse(
                stream, events=("start", "end"),
                load_dtd=False, resolve_entities=False,
                no_network=True, huge_tree=False, recover=False,
            ):
                tag = node.tag
                if not isinstance(tag, str):
                    continue
                if action == "start":
                    if tag == "CANDELA":
                        out.metadata["dtd_version"] = node.get("dtdvers")
                    if tag == "DIAGINST":
                        inside += 1
                    continue
                if tag == "QUAL" and node.getparent() is not None:
                    parent = node.getparent()
                    if parent.tag == "DIAGCLASS":
                        class_qualifiers[parent.get("id") or ""] = node.text or ""
                if tag == "DIAGINST":
                    parent = node.getparent()
                    service_index += 1
                    if service_index > 100_000:
                        raise ValueError("CDD enthält zu viele Diagnoseinstanzen.")
                    name = node.findtext("NAME/TUV") or node.findtext("QUAL") or ""
                    qualifier = node.findtext("QUAL") or ""
                    semantics = [
                        child.text.strip() for child in node.iter("SEMANTIC")
                        if child.text and child.text.strip()
                    ]
                    description = node.find("DESC")
                    desc_text = " ".join(
                        text.strip() for text in description.itertext()
                        if text and text.strip()
                    ) if description is not None else ""
                    out.records.append({
                        "category": "cdd_service",
                        "record_number": service_index,
                        "source_location": f"DIAGINST:{node.get('id', service_index)}",
                        "fin": None,
                        "reference": qualifier,
                        "label": name,
                        "data": {
                            "name": name,
                            "qualifier": qualifier,
                            "class": class_qualifiers.get(
                                parent.get("id") or "", ""
                            ) if parent is not None else "",
                            "semantics": sorted(set(semantics)),
                            "description": desc_text[:5000],
                            "service_id": node.get("id"),
                            "attributes": {
                                key: value for key, value in node.attrib.items()
                                if key in {"id", "tmplref", "req", "xauth"}
                            },
                        },
                    })
                    inside -= 1
                if inside == 0:
                    node.clear()
                    parent = node.getparent()
                    if parent is not None:
                        while node.getprevious() is not None:
                            del parent[0]
        out.metadata["service_count"] = service_index
        out.metadata["class_count"] = len(class_qualifiers)
        if not out.records:
            out.issues.append("Keine Diagnoseinstanzen in der CDD vorhanden.")
        return out


_PART_DATE = re.compile(
    r"^\s*(?P<date>\d{1,2}/\d{1,2}/\d{4}|CW\d{1,2}/\d{4})\s+"
    r"(?P<change>\S+)(?:\s+(?P<revision>Q\d{2}))?(?P<rest>.*)$",
    re.IGNORECASE,
)


def _period(label: str) -> tuple[str, str, str]:
    if label.upper().startswith("CW"):
        week, year = label[2:].split("/")
        monday = date.fromisocalendar(int(year), int(week), 1)
        return "calendar_week", monday.isoformat(), (
            monday + timedelta(days=6)
        ).isoformat()
    day = datetime.strptime(label, "%m/%d/%Y").date()
    return "day", day.isoformat(), day.isoformat()


class PartHistoryParser:
    """Read table-like production changes, retaining the original layout text."""

    name = "part_history_pdf"
    version = "1.0.0"

    def parse(self, path: Path) -> KnowledgeResult:
        from pypdf import PdfReader

        if path.stat().st_size > 35_000_000:
            raise ValueError("Part-History-PDF überschreitet die Größenbegrenzung.")
        pdf = PdfReader(path, strict=False)
        if pdf.is_encrypted or not 1 <= len(pdf.pages) <= 25:
            raise ValueError("Part-History-PDF verschlüsselt oder Seitenzahl ungültig.")
        pages = []
        result = KnowledgeResult(kind="part_history")
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text(extraction_mode="layout") or ""
            pages.append(text)
            if page_number == 1 and "Part History List" not in text:
                raise ValueError("Kein unterstütztes Part-History-Layout erkannt.")
            chunks: list[tuple[str, str]] = []
            active: list[str] = []
            for line in text.splitlines():
                match = _PART_DATE.match(line)
                if match:
                    if active:
                        chunks.append((active[0], "\n".join(active)))
                    active = [line]
                    continue
                if active:
                    if "Předlohu fomuláře" in line or "Only for internal use" in line:
                        chunks.append((active[0], "\n".join(active)))
                        active = []
                        continue
                    active.append(line)
            if active:
                chunks.append((active[0], "\n".join(active)))
            for first, raw in chunks:
                match = _PART_DATE.match(first)
                if match is None:
                    continue
                try:
                    precision, start, end = _period(match.group("date"))
                except ValueError:
                    result.issues.append("Ungültiger Datums-/Kalenderwochenwert in Part History.")
                    continue
                record_number = len(result.records) + 1
                # Descriptions can be multi-line or span multiple columns.
                # Keep raw content; do not guess an exact FIN, supplier cause,
                # introduction time or an installation population.
                result.records.append({
                    "category": "part_change",
                    "record_number": record_number,
                    "source_location": f"page:{page_number}:change:{record_number}",
                    "fin": None,
                    "reference": match.group("change"),
                    "label": match.group("revision") or match.group("change"),
                    "data": {
                        "date_label": match.group("date"),
                        "precision": precision,
                        "period_start": start,
                        "period_end": end,
                        "change_id": match.group("change"),
                        "revision": match.group("revision"),
                        "details": raw,
                        "page": page_number,
                    },
                })
        result.metadata = {
            "page_count": len(pages),
            "raw_page_texts": pages,
            "change_count": len(result.records),
            "fin_specific": False,
        }
        if not result.records:
            result.issues.append("Keine Tabellenzeilen mit Datum und Änderungsreferenz erkannt.")
        return result


def iter_autoexport_rows(path: Path) -> Iterator[dict]:
    """Yield every CSV row, including rows for FINs without a QEC case.

    Duplicate column labels receive a stable suffix rather than silently
    overwriting one another. The 1-column C1 aggregate is *not* an
    installed-population denominator for supplier failure rates.
    """
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
        reader = csv.reader(stream, delimiter=";", quotechar='"')
        header = None
        names = []
        index = 0
        for values in reader:
            if not values or not any(value.strip() for value in values):
                continue
            if header is None and (
                len(values) == 1 and (
                    values[0].startswith("Original query:")
                    or values[0].startswith("Query Id:")
                )
            ):
                continue
            if header is None:
                header = values
                seen = {}
                for col in header:
                    seen[col] = seen.get(col, 0) + 1
                    names.append(col if seen[col] == 1 else f"{col}__{seen[col]}")
                continue
            index += 1
            values += [""] * max(0, len(names) - len(values))
            row = dict(zip(names, values))
            fin = row.get("Fin") or None
            when = row.get("DiagStartDt") or row.get("RepairDt") or None
            kind = (
                "autoexport_dtc" if "EcuFaultNm" in row
                else "autoexport_repair" if "RepairDt" in row
                else "autoexport_aggregate" if names == ["C1"]
                else "autoexport_other"
            )
            yield {
                "category": kind,
                "record_number": index,
                "source_location": f"csv:row:{index}",
                "fin": fin,
                "reference": row.get("EcuFaultNm") or row.get("CntGukSid") or None,
                "label": when or "",
                "data": row,
            }


class AutoExportSummaryParser:
    """Archive a published aggregate without claiming it is a failure-rate denominator."""

    name = "autoexport_aggregate"
    version = "1.0.0"

    def parse(self, path: Path) -> KnowledgeResult:
        out = KnowledgeResult(
            kind="autoexport_summary",
            metadata={"not_an_installed_population_denominator": True},
        )
        out.records.extend(iter_autoexport_rows(path))
        if len(out.records) != 1 or out.records[0]["category"] != "autoexport_aggregate":
            raise ValueError("Kein unterstützter AutoExport-Zählwert (C1).")
        return out


class SupplierAnalysisDocumentParser:
    """Safely preserve supplier analysis text without inventing an NTF finding.

    A returned-part analysis must be matched to an explicit QEC reference
    and reviewed before its conclusion affects supplier responsibility.
    """

    name = "supplier_analysis_document"
    version = "1.0.0"

    def parse(self, path: Path) -> KnowledgeResult:
        from pypdf import PdfReader

        if path.stat().st_size > 50_000_000:
            raise ValueError("Analysebericht überschreitet die Größenbegrenzung.")
        pdf = PdfReader(path, strict=False)
        if pdf.is_encrypted or not 1 <= len(pdf.pages) <= 100:
            raise ValueError("Analysebericht verschlüsselt oder Seitenzahl ungültig.")
        first = pdf.pages[0].extract_text() or ""
        if not ("ANALYSIS" in first.upper() and "Complaint" in first):
            raise ValueError("Unbekanntes Format des Lieferanten-Analyseberichts.")
        output = KnowledgeResult(
            kind="supplier_analysis_document",
            metadata={
                "page_count": len(pdf.pages),
                "ntf_result_automatically_inferred": False,
                "requires_manual_qec_match": True,
            },
        )
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            output.records.append({
                "category": "supplier_analysis_page",
                "record_number": page_number,
                "source_location": f"pdf:page:{page_number}",
                "fin": None,  # a number in a report is not a validated case link
                "reference": None,
                "label": f"Lieferanten-Analysebericht, Seite {page_number}",
                "data": {
                    "page_number": page_number,
                    "raw_text": text,
                },
            })
        return output
