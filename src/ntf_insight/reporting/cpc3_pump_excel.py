"""Local Excel review list: CPC3 SW observations versus pump replacements.

This workbook includes FINs and must never be sent to GitHub or CI artifacts.
"""
from __future__ import annotations

from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ntf_insight.analysis.cpc3_pump_overview import CPC3PumpOverviewCase


def _safe(value: object) -> str:
    """Imported text must not become an Excel formula or external link."""
    result = str(value if value is not None else "")
    if result.lstrip().startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + result
    return result


_HEADERS = (
    "Fall-ID", "FIN", "Bauteil", "Pumpentauschdatum",
    "CPC3-SW vor Tausch", "CPC3-SW nach Tausch",
    "SW-Maßnahme / Versionswechsel", "Änderungskategorie",
    "Zeitintervall der SW-Beobachtungen",
    "CPC3-DTCs im letzten vollständigen Vortest",
    "CPC3-DTCs in vollständigen Nachtests",
    "DTCs erneut beobachtet", "DTC-Verlauf", "SW-Wirkung",
    "Vollständige CPC3-QuickTests davor",
    "Vollständige CPC3-QuickTests danach",
    "Dokumentierte SW-Maßnahmen", "Quellbeleg SW vor",
    "Quellbeleg SW nach", "Datenlücken", "AutoExport-Begrenzung",
    "Befundungsbericht prüfen", "SW-Releasenotes prüfen", "Fachliche Notiz",
)


def export_cpc3_pump_overview(
    path: str | Path,
    cases: Sequence[CPC3PumpOverviewCase],
) -> int:
    """Write the currently selected true case rows, no guessed conclusions."""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "CPC3-Pumpenfälle"
    ws.append(_HEADERS)
    for case in cases:
        ws.append((
            _safe(case.case_id),
            _safe(case.fin), _safe(case.component),
            case.swap_day.isoformat() if case.swap_day else "",
            _safe(case.before_sw),
            _safe(case.after_sw),
            _safe(case.update_status), _safe(case.update_group),
            _safe(case.interval_description),
            _safe("; ".join(case.pre_cpc3_dtcs)),
            _safe("; ".join(case.post_cpc3_dtcs)),
            _safe("; ".join(case.returned_dtcs)),
            _safe(case.dtc_status), _safe(case.sw_effect),
            case.pre_complete_tests, case.post_complete_tests,
            _safe(" | ".join(case.update_events)),
            _safe(case.before_source), _safe(case.after_source),
            _safe(case.note), "Ja" if case.truncated else "Nein",
            "Offen", "Offen", "",
        ))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:X{max(1, ws.max_row)}"
    ws.sheet_view.showGridLines = False
    widths = (22, 25, 22, 19, 24, 24, 49, 22, 27, 48, 48, 33,
              44, 57, 26, 26, 56, 60, 60, 60, 21, 26, 27, 38)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 44
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="22252A")
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="333333")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in (22, 23, 24):
            row[col - 1].font = Font(name="Arial", size=10, color="2456A6")

    info = workbook.create_sheet("Bewertungsregeln")
    info.append(("Hinweis", "Bedeutung"))
    for heading, value in (
        ("Analyseobjekt", "Eine Zeile je Pumpenfall; gleiche FIN kann mehrfach vorkommen."),
        ("Anker", "Pumpentauschdatum aus dokumentierter Quelle; nicht automatisch Ausfalldatum."),
        ("SW am Tauschtag", "Nur eine explizit dokumentierte SW-Maßnahme mit demselben Kalendertag."),
        ("SW zeitnah", "SW-Maßnahme im offengelegten Zeitfenster von 14 Tagen um den Tauschtag."),
        ("SW-Versionen", "Ein Unterschied vor/nach dem Tausch beweist kein Update am Tauschtag."),
        ("Diagnose", "Nur fallbezogene vollständige CPC3-QuickTests vor/nach dem Tauschtag."),
        ("DTC nicht erneut", "Im vorhandenen Nachtest nicht mehr gesehen; keine Heilungsbestätigung."),
        ("Ursache", "Pumpe und Software können gleichzeitig verändert worden sein; SW-Heilung ist nicht isoliert bewiesen."),
        ("Datenlücken", "Fehlende/mehrdeutige Beobachtungen bleiben explizit unbekannt."),
        ("Fachprüfung", "Softwareänderungen brauchen freigegebene Release-/Kalibrierungsunterlagen."),
        ("Datenschutz", "Enthält FINs und Werkstattbelege; nur auf freigegebenen lokalen Ablagen speichern."),
    ):
        info.append((_safe(heading), _safe(value)))
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 115
    for cell in info[1]:
        cell.fill = PatternFill("solid", fgColor="22252A")
        cell.font = Font(color="FFFFFF", bold=True)
    info.freeze_panes = "B2"
    workbook.save(path)
    return len(cases)
