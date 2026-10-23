"""Synthetische Integrationstests des lokalen KI-Analysepakets."""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from ntf_insight.domain.case import NTFCase
from ntf_insight.domain.events import PartReplacement, TipsCaseObservation
from ntf_insight.domain.vehicle import Vehicle
from ntf_insight.ingestion.artifacts import SourceType
from ntf_insight.ingestion.service import SourceImportService
from ntf_insight.reporting.ki_analysepaket import export_ki_analysepaket
from ntf_insight.storage.quicktest_store import QuickTestStore
from ntf_insight.storage.source_knowledge_store import SourceKnowledgeStore


FIN = "WDD12345678901234"
REAL_REFERENCE = "123456"
PRIVATE_TEXT = "Kunde MAX MUSTERMANN mit FIN " + FIN
SECRET_QEC_TEXT = "INTERNER BEFUND MIT FIN " + FIN


def _set_up(tmp_path):
    source = SourceImportService(tmp_path / "raw", tmp_path / "ntf-insight.duckdb")
    QuickTestStore(source.store)
    knowledge = SourceKnowledgeStore(source.store)
    swap = PartReplacement(
        vin=FIN, occurred_at=datetime(2026, 1, 5, 12),
        mileage_km=12000, component_type="Kühlmittelpumpe",
        time_precision="date",
    )
    case = NTFCase(
        case_id="QEC-FALL-VERTRAULICH",
        vehicle=Vehicle(vin=FIN, model_series="214"),
        component_type="Kühlmittelpumpe",
        component_profile_id="coolant_pump",
        replacements=[swap],
        investigated_replacement_id=swap.event_id,
        tips_observations=[
            TipsCaseObservation(
                vin=FIN, occurred_at=datetime(2026, 1, 4, 9),
                complaint=PRIVATE_TEXT, workshop_findings="Kein Fehler gesehen",
                performed_actions="Prüfstand",
            ),
        ],
        qec_reference=REAL_REFERENCE,
        qec_part_number="A-EXAMPLE-001",
    )
    source.store.save_case(case)
    records = [
        {
            "category": "qec_field",
            "source_location": "QEC/01-WE-Daten/feld:1",
            "fin": FIN,
            "reference": REAL_REFERENCE,
            "label": "Befund intern",
            "data": {
                "bereich": "WE-Daten", "name": "textbefund", "typ": "text",
                "wert": SECRET_QEC_TEXT, "beschriftung": "Befund intern",
            },
        },
        {
            "category": "qec_field",
            "source_location": "QEC/01-WE-Daten/feld:2",
            "fin": FIN,
            "reference": REAL_REFERENCE,
            "label": "Check",
            "data": {
                "bereich": "WE-Daten", "name": "check", "typ": "checkbox",
                "wert": "ja", "beschriftung": "Check",
            },
        },
    ]
    knowledge.save_records(
        "a" * 64, "qec_live_report", records,
        metadata={
            "fin": FIN, "reference": REAL_REFERENCE,
            "part_number": "A-EXAMPLE-001", "fresh_complete": True,
            "sections": 5, "fresh_sections": 5,
            "fields": 2, "values": 2, "table_rows": 0,
            "attachments": 1,
        },
    )
    with source.store.connect() as con:
        for label, tested, kilometer in (
            ("QTPRE", datetime(2026, 1, 2, 10), 11800),
            ("QTPOST", datetime(2026, 1, 10, 10), 12300),
        ):
            con.execute(
                "INSERT INTO quicktest_reports VALUES (?,?,?,?,?,?,?,?,?,?)",
                [label, FIN, FIN, tested, kilometer, 1, 1, "{}", "b" * 64, tested],
            )
            con.execute(
                "INSERT INTO quicktest_ecus VALUES (?,?,?,?,?,?,?,?,?,?)",
                [label, 1, "CPC3", "Prüfstelle", "CPC3", json.dumps({
                    "software_version": "25/04",
                    "hardware_version": "A01",
                }), 1, 1, 1, PRIVATE_TEXT],
            )
            con.execute(
                "INSERT INTO quicktest_dtcs VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                [label, 1, 1, "CPC3", "FC", "P26F100", PRIVATE_TEXT,
                 "S", 3, kilometer, kilometer, 1, 1, 1, PRIVATE_TEXT],
            )
    smr = tmp_path / "SMRD-TEST.smr-d"
    smr.write_text("SYNTHETIC_ONLY", encoding="utf-8")
    source.import_raw(smr, SourceType.SMR_D, vin_hint=FIN)
    return source


def _csv(zipfile, name):
    return list(csv.DictReader(io.StringIO(zipfile.read(name).decode("utf-8-sig"))))


def test_lightweight_export_joins_all_sources_and_retains_private_mapping(tmp_path):
    source = _set_up(tmp_path)
    target = tmp_path / "analyse.zip"
    result = export_ki_analysepaket(source, target)
    assert result.faelle == 1
    assert result.fahrzeuge == 1
    assert result.quicktests == 2
    assert result.dtcs == 2
    assert result.qec_berichte == 1
    assert not result.freitexte
    assert result.private_zuordnung.is_file()
    assert result.private_zuordnung.parent == target.parent
    assert FIN in result.private_zuordnung.read_text(encoding="utf-8-sig")
    assert REAL_REFERENCE in result.private_zuordnung.read_text(encoding="utf-8-sig")

    with ZipFile(target) as z:
        assert "Auswertung.xlsx" in z.namelist()
        assert "qec_felder.csv" in z.namelist()
        assert "quicktest_fallbezug.csv" in z.namelist()
        assert "smrd_bestand.csv" in z.namelist()
        assert all(not name.lower().endswith((".pdf", ".smr-d", ".jsonl", ".docx"))
                   for name in z.namelist())
        for name in z.namelist():
            if name.endswith(".csv") or name.endswith(".txt") or name.endswith(".json"):
                body = z.read(name).decode("utf-8-sig")
                assert FIN not in body, name
                assert REAL_REFERENCE not in body, name
                assert PRIVATE_TEXT not in body, name
                assert SECRET_QEC_TEXT not in body, name
                assert "QEC-FALL-VERTRAULICH" not in body, name
        cases = _csv(z, "faelle.csv")
        assert cases[0]["fall_id"].startswith("F")
        assert cases[0]["fahrzeug_id"].startswith("V")
        assert cases[0]["qec_bericht_status"] == "vollstaendig"
        qec = _csv(z, "qec_berichte.csv")
        assert qec[0]["fall_ids"] == cases[0]["fall_id"]
        fields = _csv(z, "qec_felder.csv")
        assert len(fields) == 2
        assert {row["wert_strukturiert"] for row in fields} == {"", "ja"}
        assert all(not row["original_wert"] for row in fields)
        rel = _csv(z, "quicktest_fallbezug.csv")
        assert {row["lage_zum_pumpentausch"] for row in rel} == {"vor_tausch", "nach_tausch"}
        before = next(r for r in rel if r["lage_zum_pumpentausch"] == "vor_tausch")
        assert before["tage_relativ_zum_tausch"] == "-3"
        dtcs = _csv(z, "dtc_beobachtungen.csv")
        assert {row["dtc_code"] for row in dtcs} == {"P26F100"}
        assert all(row["fahrzeug_id"] == cases[0]["fahrzeug_id"] for row in dtcs)
        feature = _csv(z, "fallmerkmale.csv")[0]
        cluster = _csv(z, "cluster_input.csv")[0]
        assert cluster["qt_vor_tausch"] == "1"
        assert "qt_nach_tausch" not in cluster
        assert "lieferantenbeteiligung_bewertet" not in cluster
        assert "qec_erfassungsstand" not in cluster
        assert feature["qt_vor_tausch"] == "1"
        assert feature["qt_nach_tausch"] == "1"
        assert feature["gleiche_cpc3_dtc_beidseits"] == "1"
        assert feature["smrd_bezug_bestaetigt"] == "1"
        assert _csv(z, "smrd_bestand.csv")[0]["struktur_parsebar"] == "0"
        book = load_workbook(io.BytesIO(z.read("Auswertung.xlsx")), read_only=True)
        assert set(("Übersicht", "Fallliste", "Falldaten", "Belege", "Methodik")) <= set(book.sheetnames)
        for sheet in book:
            for row in sheet.iter_rows(values_only=True):
                assert FIN not in " ".join(str(cell or "") for cell in row)
        book.close()


def test_detail_export_is_opt_in_and_not_anonymisiert(tmp_path):
    source = _set_up(tmp_path)
    target = tmp_path / "analyse-detail.zip"
    result = export_ki_analysepaket(source, target, mit_freitext=True)
    assert result.freitexte
    with ZipFile(target) as z:
        fields = _csv(z, "qec_felder.csv")
        assert SECRET_QEC_TEXT in [row["original_wert"] for row in fields]
        tips = _csv(z, "tips.csv")
        assert PRIVATE_TEXT in tips[0]["beschwerde_original"]
        manifest = json.loads(z.read("manifest.json"))
        assert manifest["fachfreitexte_enthalten"]
        assert not manifest["anonymisiert_garantiert"]


def test_existing_target_and_private_mapping_cannot_be_overwritten(tmp_path):
    source = _set_up(tmp_path)
    target = tmp_path / "paket.zip"
    target.write_bytes(b"BEHALTEN")
    with pytest.raises(FileExistsError):
        export_ki_analysepaket(source, target)
    assert target.read_bytes() == b"BEHALTEN"
    target.unlink()
    private = tmp_path / "paket_NICHT_HOCHLADEN_ZUORDNUNG.csv"
    private.write_text("GEHEIM", encoding="utf-8")
    with pytest.raises(FileExistsError):
        export_ki_analysepaket(source, target)
    assert not target.exists()
    assert private.read_text(encoding="utf-8") == "GEHEIM"


def test_no_quicktest_or_qec_still_exports_case_and_assessment_separately(tmp_path):
    source = SourceImportService(tmp_path / "raw", tmp_path / "db.duckdb")
    source.store.save_case(NTFCase(
        case_id="EXAMPLE", vehicle=Vehicle(vin=FIN), component_type="Pump",
    ))
    result = export_ki_analysepaket(source, tmp_path / "minimal.zip")
    assert result.faelle == 1
    assert result.quicktests == 0
    with ZipFile(result.pfad) as z:
        assert len(_csv(z, "faelle.csv")) == 1
        assert len(_csv(z, "dtc_beobachtungen.csv")) == 0
        assert len(_csv(z, "bewertungen_NICHT_CLUSTER.csv")) == 1



def test_ntf_workflow_label_is_not_a_clustering_feature(tmp_path):
    from ntf_insight.storage.qec_workflow_store import QecWorkflowStore

    source = _set_up(tmp_path)
    QecWorkflowStore(source.store)
    with source.store.connect() as con:
        con.execute(
            "INSERT INTO qec_workflow_decisions VALUES (?, ?, ?, ?)",
            ["c" * 64, "QEC-FALL-VERTRAULICH", "NTF", PRIVATE_TEXT],
        )
    result = export_ki_analysepaket(source, tmp_path / "labels.zip")
    with ZipFile(result.pfad) as z:
        labels = _csv(z, "bewertungen_NICHT_CLUSTER.csv")
        assert labels[0]["ntf_urteil_fachlich"] == "NTF"
        assert PRIVATE_TEXT not in z.read("bewertungen_NICHT_CLUSTER.csv").decode()
        assert "ntf_urteil" not in z.read("fallmerkmale.csv").decode()
        assert "ntf_urteil" not in z.read("cluster_input.csv").decode()
        assert FIN not in z.read("fallmerkmale.csv").decode()


def test_quicktest_only_vehicle_is_present_without_any_qec_case(tmp_path):
    source = _set_up(tmp_path)
    other_fin = "WDB00000000012345"
    with source.store.connect() as con:
        con.execute(
            "INSERT INTO quicktest_reports VALUES (?,?,?,?,?,?,?,?,?,?)",
            ["ORPHAN-QT", other_fin, other_fin, datetime(2026, 1, 7, 11),
             3000, 1, 0, "{}", "d" * 64, datetime(2026, 1, 7, 11)],
        )
    result = export_ki_analysepaket(source, tmp_path / "orphan.zip")
    assert result.fahrzeuge == 2
    assert result.quicktests == 3
    assert result.faelle == 1
    with ZipFile(result.pfad) as z:
        vehicles = _csv(z, "fahrzeuge.csv")
        assert len(vehicles) == 2
        assert any(row["anzahl_faelle"] == "0" and row["anzahl_quicktests"] == "1"
                   for row in vehicles)
        assert other_fin not in z.read("fahrzeuge.csv").decode()
