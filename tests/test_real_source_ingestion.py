"""Synthetic structural fixtures for local automotive source ZIP import.

No customer/vehicle/claim data from uploaded files is included in the repository.
"""
from __future__ import annotations

import io
import stat
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile, ZipInfo

import pytest
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from ntf_insight.ingestion.artifacts import SourceType
from ntf_insight.ingestion.batch_import import LocalBatchImporter
from ntf_insight.ingestion.real_source_parsers import (
    AutoExportParser,
    QECWorkbookParser,
    QuickTestPdfParser,
    canonical_fin,
)
from ntf_insight.ingestion.service import SourceImportService
from ntf_insight.ingestion.source_detection import detect_source_type


SYNTHETIC_FIN = "W1TTEST0000000001"


def _workbook_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Schadteile QEC"
    sheet.append([
        "Referenznr.", "FIN", "Benennung", "BR", "BM",
        "Prod.-Datum", "Rep.-Datum", "Laufleistung", "Sachnr. Teil",
        "SSL-Endbefund Benennung", "Bef.abschl.datum",
    ])
    sheet.append([
        "EXAMPLE-001", "MBC" + SYNTHETIC_FIN, "KUEHLMITTELPUMPE 100W",
        "214", "099", datetime(2025, 4, 3), datetime(2026, 1, 5),
        12345, "A000TEST", "Elektrischer Fehler - NTF",
        datetime(2026, 1, 17),
    ])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _autoexport_csv() -> bytes:
    text = (
        "Original query: synthetic fixture\n"
        "Query Id: synthetic\n\n"
        "DiagStartDt;SourceSystemNm;EcuNm;EcuHwPartnumberCd;EcuSuppHwVersionTxt;"
        "EcuSwVersionTxt;KmReadingNum;EcuFaultStateTxt;EcuFaultNm;EcuFaultTxt;"
        "FrequencyCounter;FirstOdometer;Fin\n"
        f"2026-01-02;TD;CPC3;A001;1.0;2.0;12000;S;P26F100;"
        f"Pump performance;3;11995;{SYNTHETIC_FIN}\n"
        f"2026-01-02;TD;CPC3;A001;1.0;2.0;12000;S;P260000;"
        f"Coolant circuit;1;11990;{SYNTHETIC_FIN}\n"
    )
    return text.encode("utf-8")


def _claim_csv() -> bytes:
    text = (
        "Original query: synthetic fixture\n"
        "Query Id: synthetic\n\n"
        "Fin;RepairDt;KmReadingNum;PartLetterCd;ComplDsc;RepairTypeNmGerman;"
        "WorksVisitDsc\n"
        f'{SYNTHETIC_FIN};2026-01-05;12345;A000TEST;'
        '"Pumpe funktioniert nicht";Tausch Pumpe;"Reparatur durchgeführt"\n'
    )
    return text.encode("utf-8")



def _tips_csv() -> bytes:
    """Synthetic TIPS fixture; no customer or real vehicle data."""
    return (
        "Fin;CaseId;CreationDate;InitialWorkshopMessageCustomerComplaint;"
        "InitialWorkshopMessageWorkshopFindings;"
        "InitialWorkshopMessagePerformedActions;SymptomName\n"
        f'{SYNTHETIC_FIN};TIPS-001;2026-01-03;'
        '"Kuehlmittelpumpe ohne Durchfluss";'
        '"Teststand: keine Foerderung";"Diagnose angefordert";'
        '"Kuehlmittelkreislauf"\n'
        "W1TTEST0000000002;TIPS-002;2026-01-03;"
        '"Anderes Fahrzeug";"Test";"Ohne";"Pumpe"\n'
    ).encode("utf-8")


def _quicktest_bytes(when: str, code: str) -> bytes:
    target = io.BytesIO()
    pdf = canvas.Canvas(target)
    lines = [
        f"FIN {SYNTHETIC_FIN}",
        f"VIN {SYNTHETIC_FIN}",
        "DISTANCE 12345 km",
        f"DIAG START {when}",
        "QuickTest 1",
        "ECU Count 1",
        "ECUS (1)",
        "CPC3-N127 - Control unit Drivetrain",
        "Diagnostics ID 004B35 ECU Variant CPC3_TEST",
        "Hardware",
        "Partnumber",
        "000 901 62 21 Hardware",
        "Version",
        "21/28 001",
        "Supplier Example Software",
        "Version",
        "24/25 000",
        "FAULT CODES (1)",
        f"FC {code} - Coolant pump performance. S",
        "Name Value (first occurence or unknown) Value (last occurence)",
        "frequencyCounter 2 -",
        "originalOdometerValue 12340 km -",
        "mostRecentOdometerValue 12345 km -",
    ]
    for index, line in enumerate(lines):
        # drawString preserves the simple line-wise extraction used by the parser
        pdf.drawString(35, 790 - 19 * index, line)
    pdf.save()
    return target.getvalue()


def _bundle(path: Path) -> None:
    with ZipFile(path, "w") as z:
        z.writestr("bundle/cases.xlsx", _workbook_bytes())
        z.writestr("bundle/diag.csv", _autoexport_csv())
        z.writestr("bundle/workshop.csv", _claim_csv())
        z.writestr("bundle/tips.csv", _tips_csv())
        before = _quicktest_bytes("2026-01-03 12:00:00", "P26F100")
        after = _quicktest_bytes("2026-01-10 12:00:00", "P26F100")
        z.writestr("bundle/report_pre.pdf", before)
        z.writestr("bundle/report_post.pdf", after)
        z.writestr("bundle/report_post_copy.pdf", after)
        # Same visible PDF content, different bytes/source hash.
        z.writestr(
            "bundle/report_post_metadata_variant.pdf",
            after + b"\n% synthetic producer metadata only\n",
        )
        z.writestr(
            "bundle/CPC3.cdd",
            "<?xml version='1.0'?><!DOCTYPE CANDELA SYSTEM 'candela.dtd'>"
            "<CANDELA dtdvers='20.0.5'></CANDELA>",
        )


def test_content_detection_and_cdd_not_smr_d(tmp_path) -> None:
    workbook = tmp_path / "random.xlsx"
    workbook.write_bytes(_workbook_bytes())
    diag = tmp_path / "auto.csv"
    diag.write_bytes(_autoexport_csv())
    claim = tmp_path / "claims.csv"
    claim.write_bytes(_claim_csv())
    pdf = tmp_path / "random.pdf"
    pdf.write_bytes(_quicktest_bytes("2026-01-03 12:00:00", "P26F100"))
    cdd = tmp_path / "test.cdd"
    cdd.write_text(
        "<?xml version='1.0'?><!DOCTYPE CANDELA SYSTEM 'candela.dtd'>"
        "<CANDELA></CANDELA>",
        encoding="utf-8",
    )

    assert detect_source_type(workbook) == SourceType.VEHICLE_CASE_LIST
    assert detect_source_type(diag) == SourceType.SHORT_TEST
    assert detect_source_type(claim) == SourceType.WORKSHOP_HISTORY
    tips = tmp_path / "tips.csv"
    tips.write_bytes(_tips_csv())
    assert detect_source_type(tips) == SourceType.TIPS_CASE
    assert detect_source_type(pdf) == SourceType.SHORT_TEST
    assert detect_source_type(cdd) == SourceType.DIAGNOSTIC_DESCRIPTION
    assert canonical_fin("MBC" + SYNTHETIC_FIN) == SYNTHETIC_FIN


def test_qec_parser_keeps_case_anchor_and_ntf_as_finding(tmp_path) -> None:
    path = tmp_path / "cases.xlsx"
    path.write_bytes(_workbook_bytes())
    parsed = QECWorkbookParser("a" * 64).parse(path)

    assert len(parsed.cases) == 1
    case = parsed.cases[0]
    assert case.vin == SYNTHETIC_FIN
    assert case.component_profile_id == "coolant_pump"
    assert case.investigated_replacement() is not None
    assert case.investigated_replacement().mileage_km == 12345
    assert case.investigated_replacement().time_precision == "date"
    assert case.supplier_findings[0].result.value == "NTF"


def test_csv_exports_preserve_source_row_and_datetime_precision(tmp_path) -> None:
    path = tmp_path / "diagnostics.csv"
    path.write_bytes(_autoexport_csv())
    parsed = AutoExportParser(
        "b" * 64, SourceType.SHORT_TEST, {SYNTHETIC_FIN}
    ).parse(path)
    assert len(parsed.diagnostics) == 2
    assert len(parsed.sessions) == 0  # Filtered fault export is NOT a full short test.
    assert len(parsed.ecu_states) == 1
    assert parsed.diagnostics[0].time_precision == "date"
    assert parsed.diagnostics[0].occurrence_count == 3
    assert parsed.diagnostics[0].source_ref.startswith("raw:")

    claim = tmp_path / "claims.csv"
    claim.write_bytes(_claim_csv())
    claims = AutoExportParser(
        "c" * 64, SourceType.WORKSHOP_HISTORY, {SYNTHETIC_FIN}
    ).parse(claim)
    assert len(claims.repairs) == 1
    assert "Tausch" in claims.repairs[0].action


def test_quicktest_pdf_extracts_full_session_and_dtc(tmp_path) -> None:
    path = tmp_path / "report.pdf"
    path.write_bytes(_quicktest_bytes("2026-01-03 12:00:00", "P26F100"))
    parsed = QuickTestPdfParser("d" * 64).parse(path)

    assert len(parsed.sessions) == 1
    assert len(parsed.diagnostics) == 1
    assert len(parsed.ecu_states) == 1
    assert parsed.sessions[0].dtc_count == 1
    assert parsed.diagnostics[0].ecu == "CPC3"
    assert parsed.diagnostics[0].dtc_code == "P26F100"
    assert parsed.diagnostics[0].occurrence_count == 2
    assert parsed.ecu_states[0].software_version == "24/25 000"


def test_zip_import_is_local_linked_and_idempotent(tmp_path) -> None:
    archive = tmp_path / "synthetic-sources.zip"
    _bundle(archive)
    source = SourceImportService(tmp_path / "raw", tmp_path / "cases.duckdb")
    importer = LocalBatchImporter(source)

    first = importer.import_paths([archive])
    saved = source.store.load_cases()
    assert first.cases_created == 1
    assert len(saved) == 1
    assert first.files_parsed == 8
    assert first.tips_linked == 1
    assert first.duplicates == 2
    assert len(saved[0].diagnostic_sessions) == 2
    assert len(saved[0].diagnostics) == 4
    assert len(saved[0].replacements) == 1
    assert len(saved[0].repairs) == 1
    assert len(saved[0].tips_observations) == 1
    assert saved[0].investigated_replacement() is not None
    assert source.store.load_artifacts()

    second = importer.import_paths([archive])
    repeated = source.store.load_cases()
    assert second.cases_created == 0
    assert len(repeated[0].diagnostics) == len(saved[0].diagnostics)
    assert len(repeated[0].diagnostic_sessions) == 2
    assert len(repeated[0].tips_observations) == 1


def test_zip_rejects_path_traversal_without_extracting(tmp_path) -> None:
    path = tmp_path / "unsafe.zip"
    with ZipFile(path, "w") as z:
        z.writestr("../outside.csv", b"not a valid path")
    source = SourceImportService(tmp_path / "raw", tmp_path / "cases.duckdb")
    with pytest.raises(ValueError, match="ZIP-Pfad"):
        LocalBatchImporter(source).import_paths([path])
    assert not (tmp_path / "outside.csv").exists()


def test_zip_rejects_symlinks(tmp_path) -> None:
    path = tmp_path / "symlink.zip"
    with ZipFile(path, "w") as z:
        info = ZipInfo("bundle/link.csv")
        info.create_system = 3
        info.external_attr = (stat.S_IFLNK | 0o777) << 16
        z.writestr(info, "../../secret")
    source = SourceImportService(tmp_path / "raw", tmp_path / "cases.duckdb")
    with pytest.raises(ValueError, match="symbolische Links"):
        LocalBatchImporter(source).import_paths([path])


@pytest.mark.parametrize("member_name", ["bundle/file:stream.csv", "bundle/CON.csv"])
def test_zip_rejects_windows_unsafe_member_names(tmp_path, member_name) -> None:
    path = tmp_path / "unsafe-name.zip"
    with ZipFile(path, "w") as z:
        z.writestr(member_name, b"synthetic")
    source = SourceImportService(tmp_path / "raw", tmp_path / "cases.duckdb")
    with pytest.raises(ValueError, match="ZIP"):
        LocalBatchImporter(source).import_paths([path])



def test_quicktest_and_autoexport_same_reading_counts_once_but_keeps_sources(
    tmp_path,
) -> None:
    """One QEC case, both uploads, and repeat import keep independent evidence."""
    from ntf_insight.analysis.diagnostic_observations import unique_diagnostic_observations
    from ntf_insight.analysis.dtc_clusters import build_dtc_network
    from ntf_insight.analysis.timeline import build_timeline

    qec = tmp_path / "qec.xlsx"
    qec.write_bytes(_workbook_bytes())
    quicktest = tmp_path / "quicktest.pdf"
    quicktest.write_bytes(_quicktest_bytes("2026-01-03 12:00:00", "P26F100"))
    auto = tmp_path / "diagnostics.csv"
    auto.write_bytes(
        _autoexport_csv()
        .replace(b"2026-01-02", b"2026-01-03")
        .replace(b"12000", b"12345")
    )
    source = SourceImportService(tmp_path / "raw", tmp_path / "cases.duckdb")
    importer = LocalBatchImporter(source)

    importer.import_paths([qec, quicktest, auto])
    saved = source.store.load_cases()
    assert len(saved) == 1
    case = saved[0]
    assert len(case.diagnostics) == 3  # two AutoExport rows + one QuickTest row
    observations = unique_diagnostic_observations(case.diagnostics)
    assert len(observations) == 2
    combined = next(o for o in observations if o.event.dtc_code == "P26F100")
    assert combined.has_cross_source_confirmation
    assert len(combined.source_refs) == 2
    assert len([item for item in build_timeline(case) if item.kind == "diagnostic"]) == 2
    nodes = build_dtc_network([case]).nodes
    assert next(node for node in nodes if node.key == ("CPC3", "P26F100")).event_count == 1

    importer.import_paths([qec, quicktest, auto])
    repeated = source.store.load_cases()[0]
    assert len(repeated.diagnostics) == 3
    assert len(unique_diagnostic_observations(repeated.diagnostics)) == 2



def test_qec_ssl_end_finding_without_date_is_not_backdated_to_swap(tmp_path) -> None:
    from openpyxl import load_workbook

    path = tmp_path / "undated_qec.xlsx"
    path.write_bytes(_workbook_bytes())
    book = load_workbook(path)
    book["Schadteile QEC"]["K2"] = None  # Bef.abschl.datum unknown
    book.save(path)
    book.close()

    parsed = QECWorkbookParser("f" * 64).parse(path)
    case = parsed.cases[0]
    assert case.qec_ssl_end_finding == "Elektrischer Fehler - NTF"
    assert case.qec_ssl_end_date is None
    assert case.supplier_findings == []  # no invented supplier event at swap
    assert case.investigated_replacement() is not None

    from ntf_insight.analysis.qec_findings import qec_ssl_finding_label
    assert "Befundabschlussdatum unbekannt" in qec_ssl_finding_label(case)



def test_real_autoexport_timestamp_and_duplicate_mileage_columns(tmp_path) -> None:
    """Exact UTC source wall-time is kept, not replaced with invented noon."""
    diagnostic = tmp_path / "detailed.csv"
    diagnostic.write_text(
        "DiagStartDt;DiagStartTs;Fin;EcuNm;EcuFaultNm;KmReadingNum;"
        "EcuFaultTxt;FrequencyCounter\n"
        f"{'2026-01-02'};2026-01-02T06:10:11.123Z;{SYNTHETIC_FIN};"
        "CPC3;P26F100;12345;Pump diagnosis;2\n"
        f"{'2026-01-02'};2026-01-02T07:10:11.123Z;{SYNTHETIC_FIN};"
        "CPC3;P26F100;12345;Pump diagnosis;3\n",
        encoding="utf-8",
    )
    parsed = AutoExportParser(
        "b" * 64, SourceType.SHORT_TEST, {SYNTHETIC_FIN}
    ).parse(diagnostic)
    assert parsed.processed_rows == 2
    assert parsed.invalid_rows == 0
    assert len(parsed.diagnostics) == 2
    assert parsed.diagnostics[0].occurred_at == datetime(2026, 1, 2, 6, 10, 11, 123000)
    assert getattr(parsed.diagnostics[0], "time_precision", "datetime") == "datetime"
    assert parsed.diagnostics[0].freeze_frame["autoexport_diag_start_ts"] == (
        "2026-01-02T06:10:11.123Z"
    )
    assert parsed.diagnostics[1].occurred_at.hour == 7

    claim = tmp_path / "claim_duplicate.csv"
    claim.write_text(
        "Fin;RepairDt;PartLetterCd;KmReadingNum;KmReadingNum;RepairTypeNmGerman\n"
        f"{SYNTHETIC_FIN};2026-01-05;A123;12340;12350;Tausch\n",
        encoding="utf-8",
    )
    repairs = AutoExportParser(
        "c" * 64, SourceType.WORKSHOP_HISTORY, {SYNTHETIC_FIN}
    ).parse(claim)
    assert len(repairs.repairs) == 1
    assert repairs.repairs[0].mileage_km is None
    assert "contradictory KmReadingNum" in repairs.issues[0]


def test_tips_source_and_full_vin_history_without_qec_case(tmp_path) -> None:
    """All FINs are kept in the raw knowledge store, not only QEC-linked FINs."""
    from ntf_insight.storage.source_knowledge_store import SourceKnowledgeStore

    tips = tmp_path / "tips.csv"
    tips.write_bytes(_tips_csv())
    parsed = AutoExportParser(
        "e" * 64, SourceType.TIPS_CASE, {SYNTHETIC_FIN}
    ).parse(tips)
    assert len(parsed.tips_observations) == 1
    assert parsed.skipped_rows == 1
    assert parsed.tips_observations[0].complaint == "Kuehlmittelpumpe ohne Durchfluss"
    assert parsed.tips_observations[0].performed_actions == "Diagnose angefordert"

    source = SourceImportService(tmp_path / "raw", tmp_path / "cases.duckdb")
    summary = LocalBatchImporter(source).import_paths([tips])
    store = SourceKnowledgeStore(source.store)
    assert summary.files_parsed == 1
    assert summary.tips_unassigned == 0  # no allowed QEC FINs; only raw rows
    assert store.count_rows("autoexport_tips") == 2
    assert store.fin_history_counts("W1TTEST0000000002") == {"autoexport_tips": 1}
    history = store.list_fin_history("W1TTEST0000000002")
    assert len(history) == 1
    assert history[0]["data"]["CaseId"] == "TIPS-002"


def test_tips_complaint_only_before_replacement_affects_signature(tmp_path) -> None:
    from ntf_insight.analysis.case_clustering import signature_for

    qec = tmp_path / "qec.xlsx"
    qec.write_bytes(_workbook_bytes())
    tips = tmp_path / "tips.csv"
    tips.write_bytes(_tips_csv())
    source = SourceImportService(tmp_path / "raw", tmp_path / "cases.duckdb")
    LocalBatchImporter(source).import_paths([qec, tips])
    case = source.store.load_cases()[0]
    assert len(case.tips_observations) == 1
    assert "durchfluss" in signature_for(case).complaint_tokens
    original_faults = signature_for(case).faults

    # A later workshop note may be displayed, but never leaks a diagnosis
    # conclusion or post-replacement complaint into the similarity score.
    case.tips_observations[0].occurred_at = datetime(2026, 2, 1)
    assert "durchfluss" not in signature_for(case).complaint_tokens
    assert signature_for(case).faults == original_faults


def test_fin_history_streaming_paging_and_atomic_limit(tmp_path) -> None:
    from ntf_insight.storage.duckdb_store import DuckDBStore
    from ntf_insight.storage.source_knowledge_store import SourceKnowledgeStore

    store = SourceKnowledgeStore(DuckDBStore(tmp_path / "records.duckdb"))
    records = (
        {
            "category": "autoexport_dtc",
            "fin": SYNTHETIC_FIN,
            "reference": "P26F100",
            "label": f"2026-02-{(index % 28) + 1:02d}T10:10:00Z",
            "source_location": f"csv:row:{index + 1}",
            "data": {
                "DiagStartTs": f"2026-02-{(index % 28) + 1:02d}T10:10:00Z",
                "EcuNm": "CPC3", "EcuFaultNm": "P26F100",
            },
        }
        for index in range(230)
    )
    assert store.save_records("a" * 64, "autoexport_rows", records, maximum_rows=500) == 230
    assert store.fin_history_counts(SYNTHETIC_FIN) == {"autoexport_dtc": 230}
    assert len(store.list_fin_history(SYNTHETIC_FIN, limit=200)) == 200
    assert len(store.list_fin_history(SYNTHETIC_FIN, limit=200, offset=200)) == 30
    assert store.list_fin_history(SYNTHETIC_FIN, limit=1)[0]["label"].startswith("2026-02-28")

    with pytest.raises(ValueError, match="mehr Datenzeilen"):
        store.save_records(
            "a" * 64, "autoexport_rows",
            (
                {
                    "category": "autoexport_dtc", "fin": SYNTHETIC_FIN,
                    "data": {}, "source_location": str(i),
                }
                for i in range(6)
            ),
            maximum_rows=5,
        )
    # Atomic rollback preserves all existing historic records.
    assert store.fin_history_counts(SYNTHETIC_FIN) == {"autoexport_dtc": 230}


def test_large_archive_rows_can_exceed_original_250k_limit_in_import(tmp_path, monkeypatch) -> None:
    """Verify configured capacity without storing 250k synthetic rows in CI."""
    from ntf_insight.storage.source_knowledge_store import SourceKnowledgeStore

    raw = tmp_path / "diag.csv"
    raw.write_bytes(_autoexport_csv())
    seen = []
    real_save = SourceKnowledgeStore.save_records

    def recording(self, source_sha, source_kind, records, **kwargs):
        seen.append(kwargs.get("maximum_rows"))
        return real_save(self, source_sha, source_kind, records, **kwargs)

    monkeypatch.setattr(SourceKnowledgeStore, "save_records", recording)
    source = SourceImportService(tmp_path / "raw", tmp_path / "cases.duckdb")
    result = LocalBatchImporter(source).import_paths([raw])
    assert result.reference_rows == 2
    assert seen == [1_000_000]
