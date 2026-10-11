"""Regressionstests für lokale DTC-Merkmalsfilter und Excel-Arbeitslisten."""
from __future__ import annotations

from openpyxl import load_workbook

from ntf_insight.analysis.cluster_dtc_filters import ClusterDtcFilter
from ntf_insight.analysis.cluster_representative_excel import (
    export_cluster_representatives,
)
from ntf_insight.analysis.autoexport_fin_patterns import (
    FinPatternGroup, FinPatternResult,
)


def test_ecu_dtc_filter_inclusions_exclusions_and_empty_fins():
    faults = {
        "FIN-A": (("CPC3", "P100"), ("CPC3", "P200"), ("BMS", "P100")),
        "FIN-B": (("BMS", "P200"),),
        "FIN-C": (),
    }
    assert ClusterDtcFilter().filter_signatures(faults) == {
        fin: tuple(sorted(values)) for fin, values in faults.items()
    }
    rules = ClusterDtcFilter(
        include_ecus=frozenset({"cpc3"}),
        include_faults=frozenset({("BMS", "P200")}),
        exclude_faults=frozenset({("CPC3", "P200")}),
    )
    assert rules.filter_signatures(faults) == {
        "FIN-A": (("CPC3", "P100"),),
        "FIN-B": (("BMS", "P200"),),
        "FIN-C": (),
    }
    denied = ClusterDtcFilter(
        include_ecus=frozenset({"CPC3"}),
        exclude_ecus=frozenset({"CPC3"}),
    )
    assert denied.filter_signatures(faults)["FIN-A"] == ()


def test_representative_excel_contains_only_real_cluster_fins_and_safe_text(tmp_path):
    fin = "W1TTEST0000000001"
    group = FinPatternGroup(
        group_id=4, members=((fin, 1.0), ("W1TTEST0000000002", 0.7)),
        representative_fin=fin,
        common_faults=(("CPC3", "P100", 2, "=UNTRUSTED()"),),
        qec_fin_count=1, median_similarity=0.7, representative_has_qec=True,
    )
    result = FinPatternResult(
        total_fins=3, dtc_fins=2, no_dtc_fins=("W1TTEST0000000003",),
        unassigned_fins=(), groups=(group,), code_counts={fin: 1},
        qec_links={fin: ("QEC-TEST-1",)}, initial_group_count=1,
        feature_count=1, sources_by_fin={fin: "AutoExport + QEC"},
    )
    path = tmp_path / "repr.xlsx"
    assert export_cluster_representatives(path, result) == 1
    book = load_workbook(path)
    ws = book["FIN-Repräsentanten"]
    assert ws.max_row == 2
    assert ws["A2"].value == "C4"
    assert ws["B2"].value == fin
    assert ws["F2"].value == "QEC-TEST-1"
    assert ws["K2"].value == "Offen"
    assert ws["H2"].data_type != "f"
    assert ws["I2"].number_format == "0.0%"
    assert book["Analyseinformationen"]["B7"].value.startswith(
        "Diagnoseähnlichkeit"
    )
