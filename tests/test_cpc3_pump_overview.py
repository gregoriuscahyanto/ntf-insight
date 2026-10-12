"""All-pump CPC3 SW table: synthetic cases, no real FIN or OEM SW data."""
from __future__ import annotations

import os
from datetime import datetime, timedelta

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from ntf_insight.analysis.cpc3_pump_overview import (
    build_cpc3_pump_overview, evaluate_cpc3_pump_case, is_pump_case,
)
from ntf_insight.app.widgets.cpc3_pump_overview_panel import (
    CPC3PumpOverviewPanel,
)
from ntf_insight.domain.case import NTFCase
from ntf_insight.domain.events import (
    DiagnosticEvent, DiagnosticSession, ECUStateObservation,
    PartReplacement, SoftwareEvent,
)
from ntf_insight.domain.vehicle import Vehicle
from ntf_insight.reporting.cpc3_pump_excel import export_cpc3_pump_overview


BASE = datetime(2026, 3, 1, 12)
FIN = "W1TTEST0000000001"


def _pump(
    number: int, *, vin: str = FIN, anchor: int | None = 7,
    before_sw: str = "SW-A", after_sw: str = "SW-B",
    update: int | None = None, update_action: str = "SW programmiert",
    after_codes: tuple[str, ...] = (),
    extra_post_codes: tuple[str, ...] = (),
    with_quicktests: bool = True,
    same_day_test: bool = False,
    component: str = "Coolant Pump",
) -> NTFCase:
    events = ()
    if anchor is not None:
        events = (PartReplacement(
            vin=vin, occurred_at=BASE + timedelta(days=anchor),
            time_precision="date", component_type=component,
        ),)
    states = [
        ECUStateObservation(
            vin=vin, ecu="CPC3", software_version=before_sw,
            occurred_at=BASE + timedelta(days=0),
            session_id="PRE-" + str(number),
        ),
        ECUStateObservation(
            vin=vin, ecu="CPC3", software_version=after_sw,
            occurred_at=BASE + timedelta(days=10),
            session_id="POST-" + str(number),
        ),
    ]
    sessions = []
    faults = []
    if with_quicktests:
        for name, day, codes in (
            ("PRE", 0, ("PUMP-1",)),
            ("POST", 10, after_codes),
            ("POST2", 18, extra_post_codes),
        ):
            session_id = f"{name}-{number}"
            sessions.append(DiagnosticSession(
                vin=vin, occurred_at=BASE + timedelta(days=day),
                session_id=session_id, source_is_complete_quicktest=True,
            ))
            # For a truly complete CPC3 0-DTC session, the ECU must exist
            # in the session even when there are no DiagnosticEvent rows.
            if name == "POST2":
                states.append(ECUStateObservation(
                    vin=vin, ecu="CPC3", occurred_at=BASE + timedelta(days=day),
                    software_version=after_sw, session_id=session_id,
                ))
            for code in codes:
                faults.append(DiagnosticEvent(
                    vin=vin, ecu="CPC3", dtc_code=code,
                    occurred_at=BASE + timedelta(days=day),
                    session_id=session_id,
                ))
        if same_day_test:
            sessions.append(DiagnosticSession(
                vin=vin, occurred_at=BASE + timedelta(days=anchor),
                session_id=f"SAME-{number}",
                source_is_complete_quicktest=True,
            ))
            states.append(ECUStateObservation(
                vin=vin, ecu="CPC3", software_version="SW-MIDDLE",
                occurred_at=BASE + timedelta(days=anchor),
                session_id=f"SAME-{number}",
                time_precision="date",
            ))
            faults.append(DiagnosticEvent(
                vin=vin, ecu="CPC3", dtc_code="PUMP-ON-SAME-DAY",
                session_id=f"SAME-{number}",
                occurred_at=BASE + timedelta(days=anchor),
            ))
    software_events = ()
    if update is not None:
        software_events = (SoftwareEvent(
            vin=vin, ecu="CPC3",
            occurred_at=BASE + timedelta(days=update),
            action=update_action, software_version=after_sw,
            source_ref="raw:synthetic#workshop-update",
        ),)
    return NTFCase(
        case_id=f"QEC-SYNTH-{number}",
        vehicle=Vehicle(vin=vin),
        component_type=component,
        component_profile_id="coolant_pump" if component == "Coolant Pump" else None,
        replacements=list(events), ecu_states=states,
        diagnostic_sessions=sessions, diagnostics=faults,
        software_events=list(software_events),
    )


def test_overview_keeps_multiple_qec_cases_per_fin_and_loads_source_once():
    first = _pump(1, anchor=7, update=7)
    second = _pump(2, anchor=25, update=None, after_sw="SW-A")
    other = _pump(3, vin="W1TTEST0000000002", anchor=None)
    loading = []

    def source(fin):
        loading.append(fin)
        return [], False

    report = build_cpc3_pump_overview(
        (second, first, other), load_autoexport=source,
    )
    assert report.total_pump_cases == 3
    assert report.unique_fins == 2
    assert len(loading) == 2
    assert {row.case_id for row in report.cases} == {
        "QEC-SYNTH-1", "QEC-SYNTH-2", "QEC-SYNTH-3",
    }
    assert report.documented_swap_day_updates == 1
    assert report.not_assessable >= 1


def test_documented_update_same_day_not_inferred_from_sw_version_only():
    actual = evaluate_cpc3_pump_case(_pump(1, update=7))
    assert actual.update_group == "documented_same_day"
    assert "am Tauschtag dokumentiert" in actual.update_status
    assert "SW-A" == actual.before_sw
    assert "SW-B" == actual.after_sw
    assert actual.update_events
    inferred = evaluate_cpc3_pump_case(_pump(2))
    assert inferred.update_group == "observed_change"
    assert "Updatezeit unbekannt" in inferred.update_status
    assert "dokumentiert" not in inferred.update_status


def test_same_sw_observed_is_not_claimed_as_no_update():
    row = evaluate_cpc3_pump_case(_pump(1, before_sw="SW-A", after_sw="SW-A"))
    assert row.update_group == "observed_same"
    assert "beobachtet" in row.update_status
    assert "kein Update" not in row.update_status.casefold()


def test_non_update_software_note_is_not_a_flash_document():
    row = evaluate_cpc3_pump_case(_pump(
        1, update=7, update_action="SW-Stand abgelesen",
    ))
    assert row.update_group == "observed_change"
    near = evaluate_cpc3_pump_case(_pump(2, update=10))
    assert near.update_group == "documented_nearby"
    too_far = evaluate_cpc3_pump_case(_pump(3, update=40))
    assert too_far.update_group == "observed_change"


def test_dtc_recurs_in_later_complete_qt_even_if_first_post_is_clear():
    row = evaluate_cpc3_pump_case(_pump(
        1, update=7, after_codes=(), extra_post_codes=("PUMP-1",),
        same_day_test=True,
    ))
    assert row.pre_complete_tests == 1
    assert row.post_complete_tests == 2
    assert row.returned_dtcs == ("PUMP-1",)
    assert "erneut" in row.dtc_status
    assert "nicht belegt" in row.sw_effect
    assert "PUMP-ON-SAME-DAY" not in row.post_cpc3_dtcs


def test_complete_cpc3_qt_without_post_dtc_is_observation_not_sw_cure():
    row = evaluate_cpc3_pump_case(_pump(
        1, update=7, after_codes=(), extra_post_codes=(),
    ))
    assert row.returned_dtcs == ()
    assert row.post_complete_tests == 2
    assert "nicht erneut" in row.dtc_status
    assert "nicht isolierbar" in row.sw_effect


def test_without_complete_quicktests_sw_effect_stays_unknown():
    row = evaluate_cpc3_pump_case(_pump(1, update=7, with_quicktests=False))
    assert row.pre_complete_tests == 0 and row.post_complete_tests == 0
    assert "nicht beurteilbar" in row.sw_effect
    assert "Vollständige" in row.dtc_status


def test_missing_swap_and_missing_sw_produce_explicit_data_gap():
    case = _pump(1, anchor=None, with_quicktests=False)
    case.ecu_states = []
    row = evaluate_cpc3_pump_case(case)
    assert row.update_group == "no_swap"
    assert row.swap_day is None
    assert row.before_sw is None and row.after_sw is None
    assert "nicht beurteilbar" in row.sw_effect
    assert "Pumpentauschdatum fehlt" in row.note


def test_non_pump_case_is_not_mixed_into_pump_population():
    battery = _pump(1, component="Battery")
    assert not is_pump_case(battery)
    assert build_cpc3_pump_overview([battery]).total_pump_cases == 0


def test_gui_search_filter_navigation_and_visible_excel(tmp_path):
    app = QApplication.instance() or QApplication([])
    first = _pump(1, update=7)
    second = _pump(2, update=None)
    report = build_cpc3_pump_overview([first, second])
    panel = CPC3PumpOverviewPanel([first, second], None)
    panel.show()
    app.processEvents()
    panel.set_report(report)
    assert panel.model.rowCount() == 2
    assert "2 Pumpenfälle" in panel.summary.text()
    panel.update_filter.setCurrentIndex(
        panel.update_filter.findData("documented_same_day")
    )
    assert panel.proxy.rowCount() == 1
    selected = []
    panel.case_requested.connect(selected.append)
    panel._open_case(panel.proxy.index(0, 0))
    assert selected == ["QEC-SYNTH-1"]
    filename = tmp_path / "cpc3-pumpen.xlsx"
    visible = tuple(
        panel.report.cases[
            panel.proxy.mapToSource(panel.proxy.index(i, 0)).row()
        ]
        for i in range(panel.proxy.rowCount())
    )
    assert export_cpc3_pump_overview(filename, visible) == 1
    wb = load_workbook(filename, read_only=True)
    try:
        ws = wb["CPC3-Pumpenfälle"]
        assert ws.max_row == 2
        assert ws["A2"].value == "QEC-SYNTH-1"
        assert "Bewertungsregeln" in wb.sheetnames
    finally:
        wb.close()
    panel.search.setText("does-not-exist")
    assert panel.proxy.rowCount() == 0
    panel.search.clear()
    assert panel.proxy.rowCount() == 1
    panel.close()


def test_excel_escapes_untrusted_formula_like_software_event(tmp_path):
    case = _pump(
        1, update=7, update_action="=WEBSERVICE(1)",
        before_sw="=HYPERLINK(1)", after_sw="+CMD",
    )
    row = evaluate_cpc3_pump_case(case)
    path = tmp_path / "safe.xlsx"
    export_cpc3_pump_overview(path, (row,))
    wb = load_workbook(path)
    try:
        ws = wb["CPC3-Pumpenfälle"]
        assert ws["E2"].data_type == "s"
        assert ws["E2"].value.startswith("'=")
        assert ws["F2"].data_type == "s"
        assert ws["F2"].value.startswith("'+")
    finally:
        wb.close()
